"""
Microsoft Word (.docx) Report Generator using python-docx (v2.0).
Generates executive-grade research documents with:
- Multilingual & Universal Character Support (English, Arabic, Cyrillic, Turkish, CJK, etc.)
- Automatic RTL (Right-to-Left) paragraph and run formatting via Word OpenXML (<w:bidi/>, <w:rtl/>)
- Executive cover section with metadata block
- Auto-generated Table of Contents
- Rich inline formatting (bold, italic, code as separate Word Runs)
- OpenXML custom styled callout boxes with colored border & background tint
- Blockquote rendering with gray accent bar
- Zebra-striped data tables with repeating header rows across pages
- Nested bullet support with indent levels
- Embedded images with captions
- Running headers and footers with Word page number fields (PAGE of NUMPAGES)
- Formatted code blocks with light shaded backgrounds
- Proper Word Heading 1/2/3 styles linked to Word Navigation Pane
- First-page header suppressed for clean cover page
"""

import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

RTL_REGEX = re.compile(r'[\u0590-\u05FF\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF\uFB50-\uFDFF\uFE70-\uFEFF]')

def is_rtl_text(text):
    """Detects if string contains Arabic, Hebrew, or other RTL codepoints."""
    if not isinstance(text, str):
        return False
    return bool(RTL_REGEX.search(text))

def apply_rtl_to_paragraph(p):
    """Injects Word OpenXML bidi and right alignment into paragraph properties."""
    pPr = p._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:bidi {nsdecls("w")}/>'))
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

def apply_rtl_to_run(run):
    """Injects Word OpenXML rtl and complex script font into run properties."""
    rPr = run._r.get_or_add_rPr()
    rPr.append(parse_xml(f'<w:rtl {nsdecls("w")}/>'))
    rPr.append(parse_xml(f'<w:rFonts {nsdecls("w")} w:cs="Arial"/>'))

def hex_to_rgb(hex_str):
    hex_str = hex_str.lstrip("#")
    return RGBColor(*(int(hex_str[i:i+2], 16) for i in (0, 2, 4)))


class DocxReportBuilder:
    def __init__(self, filename, title="Research Report", subtitle="", author="", organization="", date_str=""):
        self.filename = filename
        self.title = title
        self.subtitle = subtitle
        self.author = author
        self.organization = organization
        self.date_str = date_str

        self.is_rtl_doc = is_rtl_text(title) or is_rtl_text(subtitle) or is_rtl_text(author)
        self.doc = Document()

        for section in self.doc.sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)
            section.different_first_page_header_footer = True

        # Palette
        self.color_primary = hex_to_rgb("#0F172A")    # Deep Navy
        self.color_accent = hex_to_rgb("#2563EB")     # Royal Blue
        self.color_secondary = hex_to_rgb("#334155")  # Slate
        self.color_body = hex_to_rgb("#1E293B")       # Dark Charcoal
        self.color_muted = hex_to_rgb("#64748B")      # Slate-500

        self.font_family = "Arial" if self.is_rtl_doc else "Calibri"

        self._configure_styles()
        self._configure_headers_footers()

    def _configure_styles(self):
        styles = self.doc.styles

        normal = styles["Normal"]
        normal.font.name = self.font_family
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = self.color_body
        normal.paragraph_format.line_spacing = 1.15
        normal.paragraph_format.space_after = Pt(6)

        h1 = styles["Heading 1"]
        h1.font.name = self.font_family
        h1.font.size = Pt(17)
        h1.font.bold = True
        h1.font.color.rgb = self.color_primary
        h1.paragraph_format.space_before = Pt(16)
        h1.paragraph_format.space_after = Pt(6)
        h1.paragraph_format.keep_with_next = True

        h2 = styles["Heading 2"]
        h2.font.name = self.font_family
        h2.font.size = Pt(13)
        h2.font.bold = True
        h2.font.color.rgb = self.color_accent
        h2.paragraph_format.space_before = Pt(12)
        h2.paragraph_format.space_after = Pt(4)
        h2.paragraph_format.keep_with_next = True

        h3 = styles["Heading 3"]
        h3.font.name = self.font_family
        h3.font.size = Pt(11)
        h3.font.bold = True
        h3.font.color.rgb = self.color_secondary
        h3.paragraph_format.space_before = Pt(9)
        h3.paragraph_format.space_after = Pt(2)
        h3.paragraph_format.keep_with_next = True

    def _configure_headers_footers(self):
        for section in self.doc.sections:
            # --- First page header: intentionally blank for clean cover ---
            first_header = section.first_page_header
            if first_header.paragraphs:
                first_header.paragraphs[0].text = ""

            # --- First page footer: blank ---
            first_footer = section.first_page_footer
            if first_footer.paragraphs:
                first_footer.paragraphs[0].text = ""

            # --- Default header (pages 2+) ---
            header = section.header
            hp = header.paragraphs[0]
            hp.text = f"{self.title}"
            hp.style.font.size = Pt(8.5)
            hp.style.font.color.rgb = self.color_muted
            if self.is_rtl_doc:
                apply_rtl_to_paragraph(hp)
                for r in hp.runs:
                    apply_rtl_to_run(r)
            else:
                hp.alignment = WD_ALIGN_PARAGRAPH.LEFT

            # --- Default footer (pages 2+) ---
            footer = section.footer
            fp = footer.paragraphs[0]
            fp.alignment = WD_ALIGN_PARAGRAPH.LEFT if self.is_rtl_doc else WD_ALIGN_PARAGRAPH.RIGHT

            notice = "CONFIDENTIAL & PROPRIETARY — FOR AUTHORIZED USE ONLY\t\tPage "
            run_left = fp.add_run(notice)
            run_left.font.size = Pt(8.5)
            run_left.font.color.rgb = self.color_muted

            fld_page = OxmlElement('w:fldSimple')
            fld_page.set(qn('w:instr'), 'PAGE')
            fp._p.append(fld_page)

            run_mid = fp.add_run(" of ")
            run_mid.font.size = Pt(8.5)
            run_mid.font.color.rgb = self.color_muted

            fld_numpages = OxmlElement('w:fldSimple')
            fld_numpages.set(qn('w:instr'), 'NUMPAGES')
            fp._p.append(fld_numpages)

            if self.is_rtl_doc:
                apply_rtl_to_paragraph(fp)

    def build_cover_page(self):
        badge_text = "تقرير بحثي أمني وتقني استراتيجي" if self.is_rtl_doc else "STRATEGIC RESEARCH & TECHNICAL EVALUATION"
        p_badge = self.doc.add_paragraph()
        p_badge.paragraph_format.space_before = Pt(10)
        p_badge.paragraph_format.space_after = Pt(8)
        run_badge = p_badge.add_run(badge_text)
        run_badge.font.size = Pt(9.5)
        run_badge.font.bold = True
        run_badge.font.color.rgb = self.color_accent
        if self.is_rtl_doc:
            apply_rtl_to_paragraph(p_badge)
            apply_rtl_to_run(run_badge)

        p_title = self.doc.add_paragraph()
        p_title.paragraph_format.space_before = Pt(6)
        p_title.paragraph_format.space_after = Pt(8)
        run_title = p_title.add_run(self.title)
        run_title.font.size = Pt(26)
        run_title.font.bold = True
        run_title.font.color.rgb = self.color_primary
        if self.is_rtl_doc:
            apply_rtl_to_paragraph(p_title)
            apply_rtl_to_run(run_title)

        if self.subtitle:
            p_sub = self.doc.add_paragraph()
            p_sub.paragraph_format.space_before = Pt(0)
            p_sub.paragraph_format.space_after = Pt(20)
            run_sub = p_sub.add_run(self.subtitle)
            run_sub.font.size = Pt(13)
            run_sub.font.color.rgb = self.color_secondary
            if self.is_rtl_doc:
                apply_rtl_to_paragraph(p_sub)
                apply_rtl_to_run(run_sub)

        p_line = self.doc.add_paragraph()
        p_line.paragraph_format.space_after = Pt(140)
        run_bar = p_line.add_run("―" * 50)
        run_bar.font.color.rgb = hex_to_rgb("#CBD5E1")
        if self.is_rtl_doc:
            apply_rtl_to_paragraph(p_line)

        table = self.doc.add_table(rows=0, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        lbl_prepared = "إعداد الباحث" if self.is_rtl_doc else "PREPARED BY"
        lbl_org = "الجهة أو المنظمة" if self.is_rtl_doc else "ORGANIZATION"
        lbl_date = "تاريخ التقرير" if self.is_rtl_doc else "DATE OF REPORT"
        lbl_class = "مستوى السرية" if self.is_rtl_doc else "CLASSIFICATION"
        val_class = "سري للغاية / مراجعة تنفيذية" if self.is_rtl_doc else "Confidential / Executive Review"

        meta_rows = []
        if self.author:
            meta_rows.append((lbl_prepared, self.author))
        if self.organization:
            meta_rows.append((lbl_org, self.organization))
        if self.date_str:
            meta_rows.append((lbl_date, self.date_str))
        meta_rows.append((lbl_class, val_class))

        for label, val in meta_rows:
            row = table.add_row()
            cell_lbl, cell_val = (row.cells[1], row.cells[0]) if self.is_rtl_doc else (row.cells[0], row.cells[1])
            cell_lbl.width = Inches(2.0)
            cell_val.width = Inches(4.5)

            p1 = cell_lbl.paragraphs[0]
            p1.paragraph_format.space_after = Pt(4)
            r1 = p1.add_run(label)
            r1.font.bold = True
            r1.font.size = Pt(8.5)
            r1.font.color.rgb = self.color_muted

            p2 = cell_val.paragraphs[0]
            p2.paragraph_format.space_after = Pt(4)
            r2 = p2.add_run(val)
            r2.font.size = Pt(9.5)
            r2.font.color.rgb = self.color_primary

            if self.is_rtl_doc:
                apply_rtl_to_paragraph(p1)
                apply_rtl_to_run(r1)
                apply_rtl_to_paragraph(p2)
                apply_rtl_to_run(r2)

        self.doc.add_page_break()

    def add_table_of_contents(self, blocks, doc_title=""):
        """Generates a styled Table of Contents from heading blocks (H1, H2, H3)."""
        headings = [blk for blk in blocks if blk["type"] == "heading" and blk.get("level", 1) <= 3]

        # Filter out the document title heading
        if doc_title:
            headings = [h for h in headings if h["text"].strip().lower() != doc_title.strip().lower()]

        if not headings:
            return

        toc_title = "المحتويات" if self.is_rtl_doc else "TABLE OF CONTENTS"
        h = self.doc.add_heading(toc_title, level=1)
        if self.is_rtl_doc:
            apply_rtl_to_paragraph(h)
            for r in h.runs:
                apply_rtl_to_run(r)

        # Divider bar
        p_line = self.doc.add_paragraph()
        p_line.paragraph_format.space_after = Pt(10)
        run_bar = p_line.add_run("―" * 50)
        run_bar.font.color.rgb = self.color_accent
        run_bar.font.size = Pt(8)

        # TOC entries
        for entry in headings:
            level = entry.get("level", 1)
            indent = Inches(0.25 * (level - 1))
            text = entry["text"]

            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = indent
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(1)

            r = p.add_run(text)
            if level == 1:
                r.font.bold = True
                r.font.size = Pt(10.5)
                r.font.color.rgb = self.color_primary
            elif level == 2:
                r.font.size = Pt(10)
                r.font.color.rgb = self.color_accent
            else:
                r.font.size = Pt(9.5)
                r.font.color.rgb = self.color_secondary

            if is_rtl_text(text):
                apply_rtl_to_paragraph(p)
                apply_rtl_to_run(r)

        self.doc.add_page_break()

    def add_heading(self, text, level=1):
        h = self.doc.add_heading(text, level=min(level, 3))
        if is_rtl_text(text):
            apply_rtl_to_paragraph(h)
            for r in h.runs:
                apply_rtl_to_run(r)
        return h

    def add_paragraph(self, text):
        p = self.doc.add_paragraph(text)
        if is_rtl_text(text):
            apply_rtl_to_paragraph(p)
            for r in p.runs:
                apply_rtl_to_run(r)
        return p

    def add_rich_paragraph(self, runs_data):
        """Adds a paragraph with rich inline formatting (bold, italic, code).

        Args:
            runs_data: list of dicts with keys: text, bold, italic, code
        """
        p = self.doc.add_paragraph()
        rtl_detected = False

        for run_info in runs_data:
            r = p.add_run(run_info["text"])
            r.font.size = Pt(10.5)
            r.font.color.rgb = self.color_body

            if run_info.get("bold"):
                r.font.bold = True
            if run_info.get("italic"):
                r.font.italic = True
            if run_info.get("code"):
                r.font.name = "Consolas"
                r.font.size = Pt(9.5)
                r.font.color.rgb = hex_to_rgb("#0F172A")
                # Add shading to code spans
                rPr = r._r.get_or_add_rPr()
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9" w:val="clear"/>')
                rPr.append(shd)
            else:
                r.font.name = self.font_family

            if is_rtl_text(run_info["text"]):
                apply_rtl_to_run(r)
                rtl_detected = True

        if rtl_detected:
            apply_rtl_to_paragraph(p)

        return p

    def add_bullet(self, text):
        p = self.doc.add_paragraph(text, style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        if is_rtl_text(text):
            apply_rtl_to_paragraph(p)
            for r in p.runs:
                apply_rtl_to_run(r)
        return p

    def add_rich_bullet(self, runs_data, indent=0):
        """Adds a bullet with rich inline formatting and optional nesting.

        Args:
            runs_data: list of dicts with keys: text, bold, italic, code
            indent: nesting level (0, 1, 2)
        """
        p = self.doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Inches(0.25 + 0.25 * indent)
        rtl_detected = False

        for run_info in runs_data:
            r = p.add_run(run_info["text"])
            r.font.size = Pt(10.5)
            r.font.color.rgb = self.color_body

            if run_info.get("bold"):
                r.font.bold = True
            if run_info.get("italic"):
                r.font.italic = True
            if run_info.get("code"):
                r.font.name = "Consolas"
                r.font.size = Pt(9.5)
                r.font.color.rgb = hex_to_rgb("#0F172A")
                rPr = r._r.get_or_add_rPr()
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9" w:val="clear"/>')
                rPr.append(shd)
            else:
                r.font.name = self.font_family

            if is_rtl_text(run_info["text"]):
                apply_rtl_to_run(r)
                rtl_detected = True

        if rtl_detected:
            apply_rtl_to_paragraph(p)

        return p

    def add_callout(self, text, title="KEY INSIGHT", callout_type="note"):
        type_palette = {
            "note": {"border_color": "2563EB", "fill_color": "EFF6FF", "title_color": "#1D4ED8"},
            "warning": {"border_color": "D97706", "fill_color": "FFFBEB", "title_color": "#B45309"},
            "important": {"border_color": "DC2626", "fill_color": "FEF2F2", "title_color": "#B91C1C"},
            "tip": {"border_color": "059669", "fill_color": "ECFDF5", "title_color": "#047857"},
        }
        cfg = type_palette.get(callout_type, type_palette["note"])
        rtl_callout = is_rtl_text(text) or is_rtl_text(title)

        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(7.0)

        tcPr = cell._tc.get_or_add_tcPr()
        border_side = "right" if rtl_callout else "left"
        other_side = "left" if rtl_callout else "right"
        borders_elm = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'  <w:top w:val="none"/>'
            f'  <w:{border_side} w:val="single" w:sz="36" w:space="0" w:color="{cfg["border_color"]}"/>'
            f'  <w:bottom w:val="none"/>'
            f'  <w:{other_side} w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders_elm)

        shd_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{cfg["fill_color"]}"/>')
        tcPr.append(shd_elm)

        tcMar_elm = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'  <w:top w:w="140" w:type="dxa"/>'
            f'  <w:bottom w:w="140" w:type="dxa"/>'
            f'  <w:left w:w="200" w:type="dxa"/>'
            f'  <w:right w:w="200" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar_elm)

        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r_head = p.add_run(title.upper() + "\n")
        r_head.font.bold = True
        r_head.font.size = Pt(9.5)
        r_head.font.color.rgb = hex_to_rgb(cfg["title_color"])

        r_body = p.add_run(text)
        r_body.font.size = Pt(9.5)
        r_body.font.color.rgb = hex_to_rgb("#1E293B")

        if rtl_callout:
            apply_rtl_to_paragraph(p)
            apply_rtl_to_run(r_head)
            apply_rtl_to_run(r_body)

        p_space = self.doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(4)
        p_space.paragraph_format.space_after = Pt(4)

    def add_blockquote(self, text):
        """Renders a plain blockquote with a gray left accent bar and light background."""
        rtl_bq = is_rtl_text(text)

        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(7.0)

        tcPr = cell._tc.get_or_add_tcPr()
        border_side = "right" if rtl_bq else "left"
        other_side = "left" if rtl_bq else "right"
        borders_elm = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'  <w:top w:val="none"/>'
            f'  <w:{border_side} w:val="single" w:sz="24" w:space="0" w:color="94A3B8"/>'
            f'  <w:bottom w:val="none"/>'
            f'  <w:{other_side} w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders_elm)

        shd_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
        tcPr.append(shd_elm)

        tcMar_elm = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'  <w:top w:w="100" w:type="dxa"/>'
            f'  <w:bottom w:w="100" w:type="dxa"/>'
            f'  <w:left w:w="180" w:type="dxa"/>'
            f'  <w:right w:w="180" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar_elm)

        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        r.font.size = Pt(9.5)
        r.font.italic = True
        r.font.color.rgb = hex_to_rgb("#475569")

        if rtl_bq:
            apply_rtl_to_paragraph(p)
            apply_rtl_to_run(r)

        p_space = self.doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(3)
        p_space.paragraph_format.space_after = Pt(3)

    def add_image(self, path, alt=""):
        """Embeds an image with an optional caption. Gracefully skips if file not found."""
        if not os.path.exists(path):
            p = self.doc.add_paragraph()
            r = p.add_run(f"[Image not found: {alt or path}]")
            r.font.italic = True
            r.font.color.rgb = self.color_muted
            return

        try:
            self.doc.add_picture(path, width=Inches(5.5))

            if alt:
                p_caption = self.doc.add_paragraph()
                p_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_caption.paragraph_format.space_before = Pt(2)
                p_caption.paragraph_format.space_after = Pt(10)
                r_cap = p_caption.add_run(alt)
                r_cap.font.size = Pt(8.5)
                r_cap.font.italic = True
                r_cap.font.color.rgb = self.color_muted
                if is_rtl_text(alt):
                    apply_rtl_to_paragraph(p_caption)
                    apply_rtl_to_run(r_cap)
        except Exception as e:
            p = self.doc.add_paragraph()
            r = p.add_run(f"[Error loading image: {e}]")
            r.font.italic = True
            r.font.color.rgb = self.color_muted

    def add_table(self, headers, rows):
        num_rows = len(rows) + 1
        num_cols = len(headers)
        table = self.doc.add_table(rows=num_rows, cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Header Row
        hdr_cells = table.rows[0].cells
        for col_idx, text in enumerate(headers):
            cell = hdr_cells[col_idx]
            cell.text = text
            tcPr = cell._tc.get_or_add_tcPr()
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0F172A"/>')
            tcPr.append(shd)

            mar = parse_xml(
                f'<w:tcMar {nsdecls("w")}>'
                f'  <w:top w:w="120" w:type="dxa"/>'
                f'  <w:bottom w:w="120" w:type="dxa"/>'
                f'  <w:left w:w="140" w:type="dxa"/>'
                f'  <w:right w:w="140" w:type="dxa"/>'
                f'</w:tcMar>'
            )
            tcPr.append(mar)

            p = cell.paragraphs[0]
            align = WD_ALIGN_PARAGRAPH.RIGHT if is_rtl_text(text) else WD_ALIGN_PARAGRAPH.LEFT
            p.alignment = align
            if is_rtl_text(text):
                apply_rtl_to_paragraph(p)

            for run in p.runs:
                run.font.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(255, 255, 255)
                if is_rtl_text(text):
                    apply_rtl_to_run(run)

        trPr = table.rows[0]._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        # Body Rows
        for r_idx, row_data in enumerate(rows, start=1):
            row_cells = table.rows[r_idx].cells
            fill_color = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"

            trPr = table.rows[r_idx]._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

            for c_idx, val in enumerate(row_data):
                cell = row_cells[c_idx]
                val_str = str(val)
                cell.text = val_str

                tcPr = cell._tc.get_or_add_tcPr()
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
                tcPr.append(shd)

                borders = parse_xml(
                    f'<w:tcBorders {nsdecls("w")}>'
                    f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                    f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                    f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                    f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                    f'</w:tcBorders>'
                )
                tcPr.append(borders)

                mar = parse_xml(
                    f'<w:tcMar {nsdecls("w")}>'
                    f'  <w:top w:w="90" w:type="dxa"/>'
                    f'  <w:bottom w:w="90" w:type="dxa"/>'
                    f'  <w:left w:w="120" w:type="dxa"/>'
                    f'  <w:right w:w="120" w:type="dxa"/>'
                    f'</w:tcMar>'
                )
                tcPr.append(mar)

                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                align = WD_ALIGN_PARAGRAPH.RIGHT if is_rtl_text(val_str) else WD_ALIGN_PARAGRAPH.LEFT
                p.alignment = align
                if is_rtl_text(val_str):
                    apply_rtl_to_paragraph(p)

                for run in p.runs:
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = hex_to_rgb("#1E293B")
                    if is_rtl_text(val_str):
                        apply_rtl_to_run(run)

        p_space = self.doc.add_paragraph()
        p_space.paragraph_format.space_after = Pt(6)

    def add_code_block(self, code_text):
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(7.0)

        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
            f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
            f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
            f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)

        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
        tcPr.append(shd)

        mar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'  <w:top w:w="120" w:type="dxa"/>'
            f'  <w:bottom w:w="120" w:type="dxa"/>'
            f'  <w:left w:w="140" w:type="dxa"/>'
            f'  <w:right w:w="140" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(mar)

        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(code_text.strip())
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        r.font.color.rgb = hex_to_rgb("#0F172A")

        p_space = self.doc.add_paragraph()
        p_space.paragraph_format.space_after = Pt(4)

    def save(self):
        self.doc.save(self.filename)
        return self.filename
